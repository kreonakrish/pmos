"""Parse binary file formats into plain text for RAG ingestion."""

from __future__ import annotations

import csv
import io
import json
from typing import Optional

from app.utils.logger import get_logger

logger = get_logger(layer="file_parser")


def parse_file(content_bytes: bytes, filename: str, mime_type: str = "") -> str:
    """Extract text from a file based on its extension or MIME type.

    Returns plain text suitable for chunking and embedding.
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    try:
        if ext == "pdf" or "pdf" in mime_type:
            return _parse_pdf(content_bytes)
        elif ext == "docx" or "wordprocessingml" in mime_type:
            return _parse_docx(content_bytes)
        elif ext == "xlsx" or "spreadsheetml" in mime_type:
            return _parse_xlsx(content_bytes)
        elif ext in ("html", "htm") or "html" in mime_type:
            return _parse_html(content_bytes)
        elif ext == "csv" or "csv" in mime_type:
            return _parse_csv(content_bytes)
        elif ext == "json" or "json" in mime_type:
            return _parse_json(content_bytes)
        else:
            # Plain text fallback (txt, md, etc.)
            return content_bytes.decode("utf-8", errors="replace")
    except Exception as exc:
        logger.error("File parse failed, falling back to raw text", filename=filename, error=str(exc))
        return content_bytes.decode("utf-8", errors="replace")


def _parse_pdf(data: bytes) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            pages.append(f"--- Page {i + 1} ---\n{text.strip()}")
    result = "\n\n".join(pages)
    logger.info("PDF parsed", pages=len(reader.pages), chars=len(result))
    return result


def _parse_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    # Also extract tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    result = "\n\n".join(paragraphs)
    logger.info("DOCX parsed", paragraphs=len(paragraphs), chars=len(result))
    return result


def _parse_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c for c in cells):
                rows.append(",".join(cells))
        if rows:
            sheets.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows))
    wb.close()
    result = "\n\n".join(sheets)
    logger.info("XLSX parsed", sheets=len(sheets), chars=len(result))
    return result


def _parse_html(data: bytes) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "html.parser")
    # Remove script/style elements
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    text = soup.get_text(separator="\n", strip=True)
    logger.info("HTML parsed", chars=len(text))
    return text


def _parse_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = []
    for row in reader:
        rows.append(",".join(row))
    result = "\n".join(rows)
    logger.info("CSV parsed", rows=len(rows), chars=len(result))
    return result


def _parse_json(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    try:
        obj = json.loads(text)
        return json.dumps(obj, indent=2, default=str)
    except json.JSONDecodeError:
        return text
