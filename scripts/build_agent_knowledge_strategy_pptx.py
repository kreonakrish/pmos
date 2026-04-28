"""Generate the 'Skills, Scripts, and RAG' agent-knowledge-strategy deck.

A pragmatic engineering brief comparing three ways to give AI agents
domain-specific knowledge: dedicated scripts, skills (markdown instruction
files), and RAG knowledge bases. Covers tradeoffs, when to use each,
staleness behavior, and a recommended layered architecture.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import List, Sequence, Tuple

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


# ── Theme tokens ──────────────────────────────────────────────────────────────
BG = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x2A, 0x44)
INK_MUTED = RGBColor(0x55, 0x60, 0x7A)
ACCENT = RGBColor(0x1B, 0x55, 0xE0)
ACCENT_DARK = RGBColor(0x14, 0x40, 0xA8)
WARN = RGBColor(0xC1, 0x46, 0x2C)
GOOD = RGBColor(0x1F, 0x80, 0x55)
RULE = RGBColor(0xE3, 0xE7, 0xEF)
HEADER_BG = RGBColor(0x1F, 0x2A, 0x44)
HEADER_FG = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = RGBColor(0xF5, 0xF7, 0xFB)

# Per-approach accent colors (used on column headers / chips)
COLOR_SCRIPT = RGBColor(0x1F, 0x80, 0x55)   # green — deterministic
COLOR_SKILL = RGBColor(0xC2, 0x6E, 0x1A)    # amber — judgment under uncertainty
COLOR_RAG = RGBColor(0x1B, 0x55, 0xE0)      # blue — recall over corpus
COLOR_SCRIPT_BG = RGBColor(0xE9, 0xF5, 0xEE)
COLOR_SKILL_BG = RGBColor(0xFB, 0xF1, 0xE2)
COLOR_RAG_BG = RGBColor(0xEC, 0xF3, 0xFF)


# ── Helpers ───────────────────────────────────────────────────────────────────
def add_blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_slide_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, *,
                size=18, bold=False, color=INK, align=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.TOP, font_name="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    if isinstance(text, list):
        first = True
        for t in text:
            if first:
                run = p.add_run()
                first = False
            else:
                p2 = tf.add_paragraph()
                p2.alignment = align
                run = p2.add_run()
            run.text = t
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            run.font.name = font_name
    else:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font_name
    return tb


def add_bullets(slide, left, top, width, height, items: List[Tuple[str, int]],
                *, size=14, color=INK, bullet_color=None):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    for i, (text, level) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = level
        bullet = "•" if level == 0 else "–"
        run = p.add_run()
        run.text = f"{bullet}  {text}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"
        p.space_after = Pt(5)
    return tb


def add_header_bar(slide, title, subtitle=None, *, slide_num=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(0), Inches(0), Inches(13.333), Inches(0.18))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    add_textbox(slide, Inches(0.55), Inches(0.32), Inches(11), Inches(0.6),
                title, size=28, bold=True, color=INK)
    if subtitle:
        add_textbox(slide, Inches(0.55), Inches(0.85), Inches(11), Inches(0.4),
                    subtitle, size=14, color=INK_MUTED)
    if slide_num is not None:
        add_textbox(slide, Inches(12.5), Inches(7.05), Inches(0.7), Inches(0.3),
                    str(slide_num), size=10, color=INK_MUTED, align=PP_ALIGN.RIGHT)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                  Inches(0.55), Inches(7.0), Inches(12.2), Inches(0.02))
    rule.fill.solid(); rule.fill.fore_color.rgb = RULE
    rule.line.fill.background()


def add_table(slide, left, top, width, height,
              headers: Sequence[str], rows: Sequence[Sequence[str]],
              *, header_bg=HEADER_BG, header_fg=HEADER_FG,
              col_widths_in: Sequence[float] = None,
              cell_size=12,
              header_colors: Sequence[RGBColor] = None):
    n_cols = len(headers)
    n_rows = len(rows) + 1
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    tbl = tbl_shape.table
    if col_widths_in:
        for i, w in enumerate(col_widths_in):
            tbl.columns[i].width = Inches(w)
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        bg = header_colors[j] if header_colors else header_bg
        cell.fill.solid(); cell.fill.fore_color.rgb = bg
        tf = cell.text_frame
        tf.margin_left = Emu(60000); tf.margin_right = Emu(60000)
        tf.margin_top = Emu(40000); tf.margin_bottom = Emu(40000)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        run = p.add_run(); run.text = h
        run.font.size = Pt(cell_size + 1); run.font.bold = True
        run.font.color.rgb = header_fg; run.font.name = "Calibri"
    for i, row in enumerate(rows):
        bg = BG if i % 2 == 0 else ROW_ALT
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            cell.fill.solid(); cell.fill.fore_color.rgb = bg
            tf = cell.text_frame
            tf.margin_left = Emu(60000); tf.margin_right = Emu(60000)
            tf.margin_top = Emu(30000); tf.margin_bottom = Emu(30000)
            tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            run = p.add_run(); run.text = str(val)
            run.font.size = Pt(cell_size); run.font.color.rgb = INK
            run.font.name = "Calibri"
    return tbl_shape


def add_chip(slide, left, top, text, *, color=ACCENT, fg=RGBColor(0xFF, 0xFF, 0xFF),
             size=11, width=None, height=Inches(0.32)):
    if width is None:
        width = Inches(max(1.0, 0.13 * len(text)))
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    chip.fill.solid(); chip.fill.fore_color.rgb = color
    chip.line.fill.background()
    tf = chip.text_frame
    tf.margin_left = Emu(80000); tf.margin_right = Emu(80000)
    tf.margin_top = Emu(20000); tf.margin_bottom = Emu(20000)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = text
    run.font.size = Pt(size); run.font.bold = True
    run.font.color.rgb = fg; run.font.name = "Calibri"
    return chip


def add_panel(slide, left, top, width, height, *,
              title: str, body_lines: List[str],
              accent: RGBColor, bg: RGBColor,
              title_size=15, body_size=12):
    """A coloured-edge panel with a heading and bullet-style lines."""
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid(); box.fill.fore_color.rgb = bg
    box.line.color.rgb = accent
    box.line.width = Pt(1.25)
    add_textbox(slide, left + Inches(0.18), top + Inches(0.12),
                width - Inches(0.3), Inches(0.4),
                title, size=title_size, bold=True, color=accent)
    add_bullets(slide, left + Inches(0.18), top + Inches(0.55),
                width - Inches(0.3), height - Inches(0.7),
                [(line, 0) for line in body_lines],
                size=body_size, color=INK)


# ── Slides ────────────────────────────────────────────────────────────────────
def slide_title(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                              Inches(0), Inches(0), Inches(13.333), Inches(3.0))
    band.fill.solid(); band.fill.fore_color.rgb = ACCENT_DARK
    band.line.fill.background()
    add_textbox(s, Inches(0.7), Inches(0.7), Inches(11.5), Inches(0.5),
                "ENGINEERING BRIEF", size=14, bold=True,
                color=RGBColor(0xCB, 0xD8, 0xF7))
    add_textbox(s, Inches(0.7), Inches(1.2), Inches(11.5), Inches(1.4),
                "Skills, Scripts, and RAG", size=44, bold=True,
                color=RGBColor(0xFF, 0xFF, 0xFF))
    add_textbox(s, Inches(0.7), Inches(2.2), Inches(11.5), Inches(0.6),
                "Choosing the right knowledge surface for AI agents",
                size=22, bold=True,
                color=RGBColor(0xCB, 0xD8, 0xF7))
    add_textbox(s, Inches(0.7), Inches(3.5), Inches(11.5), Inches(0.5),
                "Three tools. Three shapes of knowledge. One pragmatic rule.",
                size=22, bold=True, color=INK)
    add_textbox(s, Inches(0.7), Inches(4.1), Inches(11.5), Inches(0.5),
                "When to write a script, when to write a skill, when to reach for RAG.",
                size=18, color=INK_MUTED)
    add_textbox(s, Inches(0.7), Inches(6.7), Inches(11.5), Inches(0.4),
                "Internal · For engineering review",
                size=11, color=INK_MUTED)


def slide_problem(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    add_header_bar(s, "The Problem We Are Trying to Solve",
                   "Teams ship skills like Confluence pages — and the bill comes due later.",
                   slide_num=2)
    add_textbox(s, Inches(0.55), Inches(1.45), Inches(12.2), Inches(0.4),
                "What we keep seeing", size=16, bold=True, color=ACCENT_DARK)
    add_bullets(s, Inches(0.7), Inches(1.85), Inches(12.0), Inches(2.0), [
        ("Skills proliferate as the default unit of \"sharing knowledge with the agent.\"", 0),
        ("They behave like Confluence docs: easy to write, easy to forget, hard to verify.", 0),
        ("Drift is silent — the code changes, the markdown does not, and the agent confidently follows the stale instructions.", 0),
        ("Retrieval is probabilistic — the agent may pick the wrong skill, or none at all, on a given query.", 0),
    ], size=14)
    add_textbox(s, Inches(0.55), Inches(4.3), Inches(12.2), Inches(0.4),
                "What this brief argues", size=16, bold=True, color=ACCENT_DARK)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(0.55), Inches(4.75), Inches(12.2), Inches(1.95))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xEC, 0xF3, 0xFF)
    box.line.color.rgb = ACCENT; box.line.width = Pt(1.0)
    add_textbox(s, Inches(0.85), Inches(4.9), Inches(11.6), Inches(0.4),
                "Skills are not wrong — they are misused.",
                size=15, bold=True, color=ACCENT_DARK)
    add_bullets(s, Inches(0.85), Inches(5.35), Inches(11.6), Inches(1.4), [
        ("Most \"skill-shaped\" needs are actually script-shaped — promote them only when judgment is required.", 0),
        ("RAG is the most probabilistic of the three, not the least. It is not the staleness fix.", 0),
        ("The honest answer is layered: scripts where you can, skills where you must, RAG only when corpus size or volatility forces it.", 0),
    ], size=13)


def slide_at_a_glance(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    add_header_bar(s, "Three Approaches at a Glance",
                   "Same goal — giving the agent domain knowledge. Different shapes of knowledge.",
                   slide_num=3)
    headers = ["Approach", "Shape of knowledge", "Determinism", "Best for"]
    rows = [
        ["Script (executable)",
         "Sequence of steps with fixed inputs / outputs",
         "Deterministic",
         "Operations the team already knows how to do"],
        ["Skill (instruction file)",
         "Pattern of decisions, composable building blocks",
         "Probabilistic at selection — fixed once selected",
         "Tasks needing judgment + composition"],
        ["RAG / knowledge base",
         "Large corpus indexed for similarity recall",
         "Probabilistic at retrieval (most uncertain of the three)",
         "Volatile or volume-heavy reference material"],
    ]
    add_table(s, Inches(0.55), Inches(1.6), Inches(12.2), Inches(3.6),
              headers, rows, col_widths_in=[2.6, 4.0, 2.6, 3.0],
              cell_size=12,
              header_colors=[HEADER_BG, COLOR_SCRIPT, COLOR_SKILL, COLOR_RAG])
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(0.55), Inches(5.6), Inches(12.2), Inches(1.0))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xEC, 0xF3, 0xFF)
    box.line.color.rgb = ACCENT; box.line.width = Pt(1.0)
    add_textbox(s, Inches(0.85), Inches(5.75), Inches(11.6), Inches(0.4),
                "The lens that resolves the rest of the deck",
                size=14, bold=True, color=ACCENT_DARK)
    add_textbox(s, Inches(0.85), Inches(6.1), Inches(11.6), Inches(0.5),
                "If steps are enumerable → script. If composition under judgment → skill. If too big or too volatile to encode → RAG.",
                size=13, color=INK)


def slide_script_deepdive(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    add_header_bar(s, "Approach 1 — Scripts",
                   "An executable artifact is the cleanest possible specification of \"how to do X.\"",
                   slide_num=4)
    add_chip(s, Inches(0.55), Inches(1.45), "DETERMINISTIC", color=COLOR_SCRIPT)
    add_panel(s, Inches(0.55), Inches(2.0), Inches(6.0), Inches(2.6),
              title="Strengths",
              body_lines=[
                  "Fully deterministic — same inputs, same outputs.",
                  "Drift is loud: the script breaks when the world changes.",
                  "Testable: unit + integration tests gate every change.",
                  "Single source of truth — runnable, not just readable.",
                  "Reviewable in code review — diffs are concrete.",
              ],
              accent=COLOR_SCRIPT, bg=COLOR_SCRIPT_BG)
    add_panel(s, Inches(6.75), Inches(2.0), Inches(6.0), Inches(2.6),
              title="Weaknesses",
              body_lines=[
                  "Rigid — 90% of cases solved, 10% fail or need code change.",
                  "No judgment under uncertainty.",
                  "Authoring cost is real engineer-time, not minutes.",
                  "Skill ceiling is the author's foresight at write-time.",
              ],
              accent=WARN, bg=RGBColor(0xFE, 0xF6, 0xE6))
    add_textbox(s, Inches(0.55), Inches(4.85), Inches(12.2), Inches(0.4),
                "Use when…", size=15, bold=True, color=COLOR_SCRIPT)
    add_bullets(s, Inches(0.7), Inches(5.25), Inches(12.0), Inches(1.5), [
        ("The steps are known and stable: deploys, migrations, seeds, validations, regression harnesses.", 0),
        ("The cost of getting it wrong is high: drop tables, rotate keys, push to prod.", 0),
        ("You want drift to surface as a build/test failure rather than a silent agent mistake.", 0),
    ], size=13)
    add_textbox(s, Inches(0.55), Inches(6.6), Inches(12.2), Inches(0.4),
                "Examples in PMOS today: scripts/seed_tool_datasource_links.py · scripts/cross_system_query.py · scripts/deploy.sh",
                size=11, color=INK_MUTED)


def slide_skill_deepdive(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    add_header_bar(s, "Approach 2 — Skills",
                   "Compressed judgment for tasks that resist deterministic encoding.",
                   slide_num=5)
    add_chip(s, Inches(0.55), Inches(1.45), "JUDGMENT UNDER UNCERTAINTY", color=COLOR_SKILL)
    add_panel(s, Inches(0.55), Inches(2.0), Inches(6.0), Inches(2.6),
              title="Strengths",
              body_lines=[
                  "Composable — agent picks patterns and recombines.",
                  "Cheap to author — markdown, not code.",
                  "Lives in a fixed slot once invoked (less probabilistic than RAG).",
                  "Right home for \"how to think about X here\" knowledge.",
              ],
              accent=COLOR_SKILL, bg=COLOR_SKILL_BG)
    add_panel(s, Inches(6.75), Inches(2.0), Inches(6.0), Inches(2.6),
              title="Weaknesses",
              body_lines=[
                  "Silent staleness — no test fails when the system drifts.",
                  "Selection is probabilistic — agent may not pick the right skill.",
                  "Quality scales with curation discipline; sprawl is the default failure mode.",
                  "No execution guarantees — instructions, not invariants.",
              ],
              accent=WARN, bg=RGBColor(0xFE, 0xF6, 0xE6))
    add_textbox(s, Inches(0.55), Inches(4.85), Inches(12.2), Inches(0.4),
                "Use when…", size=15, bold=True, color=COLOR_SKILL)
    add_bullets(s, Inches(0.7), Inches(5.25), Inches(12.0), Inches(1.5), [
        ("The work requires composition under judgment, not a sequence — code review, refactor scoping, design tradeoffs.", 0),
        ("You can link the skill to source-of-truth files and add a test that fails if those files move.", 0),
        ("The audience is the agent itself, not a human reader (otherwise it is just a doc).", 0),
    ], size=13)
    add_textbox(s, Inches(0.55), Inches(6.6), Inches(12.2), Inches(0.4),
                "Treat skills like code: version with the codebase, link to live paths, write tests that catch drift.",
                size=11, bold=True, color=INK_MUTED)


def slide_rag_deepdive(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    add_header_bar(s, "Approach 3 — RAG / Knowledge Base",
                   "Recall over a corpus too big to fit in context and too volatile to hard-code.",
                   slide_num=6)
    add_chip(s, Inches(0.55), Inches(1.45), "VOLATILE / HIGH-VOLUME RECALL", color=COLOR_RAG)
    add_panel(s, Inches(0.55), Inches(2.0), Inches(6.0), Inches(2.6),
              title="Strengths",
              body_lines=[
                  "Scales to large corpora the model cannot hold in context.",
                  "Indexer can refresh from live sources — freshness by design.",
                  "Single retrieval surface across many document types.",
                  "Decouples knowledge growth from prompt engineering.",
              ],
              accent=COLOR_RAG, bg=COLOR_RAG_BG)
    add_panel(s, Inches(6.75), Inches(2.0), Inches(6.0), Inches(2.6),
              title="Weaknesses",
              body_lines=[
                  "Most probabilistic of the three — retrieval ranking can miss.",
                  "Wrong-chunk failures are silent and hard to detect downstream.",
                  "Infra cost: index, embed, refresh, evaluate — non-trivial.",
                  "Garbage-in is amplified — a stale corpus degrades every answer.",
              ],
              accent=WARN, bg=RGBColor(0xFE, 0xF6, 0xE6))
    add_textbox(s, Inches(0.55), Inches(4.85), Inches(12.2), Inches(0.4),
                "Use when…", size=15, bold=True, color=COLOR_RAG)
    add_bullets(s, Inches(0.7), Inches(5.25), Inches(12.0), Inches(1.5), [
        ("The knowledge surface is too large to encode statically: ADRs, runbooks, ontology, customer history.", 0),
        ("The content changes faster than humans can re-curate skills or scripts.", 0),
        ("You can wire freshness checks and retrieval-evaluation into CI — otherwise it rots like everything else.", 0),
    ], size=13)
    add_textbox(s, Inches(0.55), Inches(6.6), Inches(12.2), Inches(0.4),
                "Counter-intuitive: choosing RAG to fix \"probabilistic skill retrieval\" usually makes things more probabilistic, not less.",
                size=11, bold=True, color=INK_MUTED)


def slide_decision_matrix(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    add_header_bar(s, "Decision Matrix — How to Choose",
                   "A four-question gate. Default to script unless the higher-cost option earns its keep.",
                   slide_num=7)
    headers = ["Question", "If yes →", "If no →"]
    rows = [
        ["1. Can the steps be enumerated as code with stable inputs and outputs?",
         "Write a script. Stop here.",
         "Continue."],
        ["2. Does the task require composition or judgment that varies by situation?",
         "Write a skill. Link to source-of-truth files; add a drift test.",
         "Continue."],
        ["3. Is the relevant corpus too large to fit in context or too volatile to encode?",
         "Reach for RAG. Wire in freshness + retrieval evaluation.",
         "Reconsider — most things land in script or skill."],
        ["4. Is the answer changing in real time (live state, current incidents)?",
         "Use a queryable system (catalog, graph, DB) — RAG only if no API.",
         "Stop. Pick from the previous answers."],
    ]
    add_table(s, Inches(0.55), Inches(1.6), Inches(12.2), Inches(4.4),
              headers, rows, col_widths_in=[5.2, 3.7, 3.3], cell_size=12)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(0.55), Inches(6.15), Inches(12.2), Inches(0.65))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xEC, 0xF3, 0xFF)
    box.line.color.rgb = ACCENT; box.line.width = Pt(1.0)
    add_textbox(s, Inches(0.85), Inches(6.25), Inches(11.6), Inches(0.5),
                "Default rule: if you can write a script, write the script. The other two earn their keep.",
                size=14, bold=True, color=ACCENT_DARK)


def slide_staleness(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    add_header_bar(s, "Staleness — Who Notices First?",
                   "How each approach behaves when the codebase changes underneath it.",
                   slide_num=8)
    headers = ["Approach", "Drift behavior", "Time to discover", "Mitigation"]
    rows = [
        ["Script",
         "Breaks loudly — exit code, test failure, type error.",
         "Minutes (next CI run).",
         "Already covered by your test pyramid."],
        ["Skill",
         "Silent rot. Agent follows stale instructions and produces a confidently-wrong action.",
         "Days to weeks — discovered via incident.",
         "Link to file paths in code; CI test asserts the paths still exist."],
        ["RAG",
         "Silent degradation. Retrieval picks an outdated chunk; answer drifts.",
         "Weeks to months — usually noticed via aggregate quality dip.",
         "Indexer freshness SLO + retrieval evaluation harness on a fixed eval set."],
    ]
    add_table(s, Inches(0.55), Inches(1.6), Inches(12.2), Inches(3.6),
              headers, rows,
              col_widths_in=[2.0, 4.4, 2.4, 3.4],
              cell_size=12,
              header_colors=[HEADER_BG, COLOR_SCRIPT, COLOR_SKILL, COLOR_RAG])
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(0.55), Inches(5.6), Inches(12.2), Inches(1.1))
    box.fill.solid(); box.fill.fore_color.rgb = COLOR_SCRIPT_BG
    box.line.color.rgb = COLOR_SCRIPT; box.line.width = Pt(1.0)
    add_textbox(s, Inches(0.85), Inches(5.75), Inches(11.6), Inches(0.4),
                "Why scripts win the staleness game",
                size=14, bold=True, color=COLOR_SCRIPT)
    add_textbox(s, Inches(0.85), Inches(6.15), Inches(11.6), Inches(0.5),
                "Loud failure beats silent rot. A script that breaks in CI is a feature, not a bug — drift becomes a build error, not a production incident.",
                size=13, color=INK)


def slide_probabilistic(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    add_header_bar(s, "Determinism vs. Probabilistic Retrieval",
                   "The honest ranking — and why \"reach for RAG to fix probabilistic skills\" is the wrong move.",
                   slide_num=9)
    headers = ["Approach", "Where probability enters", "What is fixed once chosen"]
    rows = [
        ["Script",
         "Nowhere — fully deterministic.",
         "Everything."],
        ["Skill",
         "Selection: agent decides whether to invoke this skill on a given query.",
         "Once invoked, the skill content is fixed in the system prompt."],
        ["RAG",
         "Two stages: query rewrite → retrieval ranking. Both can miss.",
         "Only the chunks actually retrieved make it into context — and they may be wrong."],
    ]
    add_table(s, Inches(0.55), Inches(1.6), Inches(12.2), Inches(3.0),
              headers, rows,
              col_widths_in=[2.0, 5.6, 4.6], cell_size=12,
              header_colors=[HEADER_BG, COLOR_SCRIPT, COLOR_SKILL])
    add_textbox(s, Inches(0.55), Inches(5.0), Inches(12.2), Inches(0.4),
                "The counter-intuitive takeaway", size=16, bold=True, color=ACCENT_DARK)
    add_bullets(s, Inches(0.7), Inches(5.4), Inches(12.0), Inches(1.5), [
        ("RAG is more probabilistic than skills, not less. Two ranking decisions stack on top of each other.", 0),
        ("Skills sit in a fixed slot once invoked — RAG chunks are sampled per-query.", 0),
        ("If your concern is \"agent picks the wrong knowledge,\" RAG is not the fix; tighter scoping + fewer skills is.", 0),
    ], size=13)


def slide_layered(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    add_header_bar(s, "Recommended Layered Architecture",
                   "All three at once — used for what they are each best at.",
                   slide_num=10)
    add_panel(s, Inches(0.55), Inches(1.55), Inches(4.05), Inches(4.4),
              title="Operational layer",
              body_lines=[
                  "Scripts you commit and run.",
                  "Deploys, migrations, seeds.",
                  "Validation harnesses.",
                  "Cron / scheduled jobs.",
                  "—",
                  "Drift surface: CI / tests.",
                  "Owner: the engineer.",
                  "Default tool: code.",
              ],
              accent=COLOR_SCRIPT, bg=COLOR_SCRIPT_BG,
              title_size=15, body_size=12)
    add_panel(s, Inches(4.7), Inches(1.55), Inches(4.05), Inches(4.4),
              title="Judgment layer",
              body_lines=[
                  "Skills the agent invokes.",
                  "Code review patterns.",
                  "Refactor scoping rubrics.",
                  "Design-tradeoff playbooks.",
                  "—",
                  "Drift surface: linked-path tests.",
                  "Owner: the team that writes them.",
                  "Default tool: markdown + tests.",
              ],
              accent=COLOR_SKILL, bg=COLOR_SKILL_BG,
              title_size=15, body_size=12)
    add_panel(s, Inches(8.85), Inches(1.55), Inches(3.9), Inches(4.4),
              title="Volatile-knowledge layer",
              body_lines=[
                  "RAG / knowledge graph / catalog.",
                  "ADRs, runbooks, incidents.",
                  "Live ontology + bindings.",
                  "Customer / domain corpora.",
                  "—",
                  "Drift surface: indexer SLO + eval set.",
                  "Owner: a data team.",
                  "Default tool: query, not encode.",
              ],
              accent=COLOR_RAG, bg=COLOR_RAG_BG,
              title_size=15, body_size=12)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(0.55), Inches(6.15), Inches(12.2), Inches(0.65))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xEC, 0xF3, 0xFF)
    box.line.color.rgb = ACCENT; box.line.width = Pt(1.0)
    add_textbox(s, Inches(0.85), Inches(6.25), Inches(11.6), Inches(0.5),
                "PMOS today already does this — scripts in /scripts, ontology in Neo4j, RAG in Qdrant. The trap is letting the skill layer sprawl into the others.",
                size=13, bold=True, color=ACCENT_DARK)


def slide_antipatterns(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    add_header_bar(s, "Anti-Patterns to Avoid",
                   "The four ways teams over-spend on the wrong layer.",
                   slide_num=11)
    headers = ["Anti-pattern", "What it looks like", "What to do instead"]
    rows = [
        ["Skill-for-everything",
         "A markdown for \"how to deploy,\" \"how to run tests,\" \"how to seed data.\"",
         "Write the script. The skill, if needed, just says \"run scripts/X.\""],
        ["Confluence-as-a-skill",
         "Documentation ported into a /skills folder — no decision authority, no actions.",
         "Leave it as documentation. Skills are for the agent, not for humans."],
        ["RAG over the source code",
         "Indexing the repo to answer code questions — the IDE already does this better.",
         "Use the IDE / language server. Reserve RAG for non-code corpora."],
        ["Script-for-judgment",
         "review.py that grades code with regexes; estimate.py that grades scope by file count.",
         "Promote to a skill — the agent can apply judgment that regex cannot."],
    ]
    add_table(s, Inches(0.55), Inches(1.6), Inches(12.2), Inches(4.6),
              headers, rows,
              col_widths_in=[2.6, 5.4, 4.2], cell_size=12)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(0.55), Inches(6.4), Inches(12.2), Inches(0.45))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xFE, 0xF6, 0xE6)
    box.line.color.rgb = WARN; box.line.width = Pt(1.0)
    add_textbox(s, Inches(0.85), Inches(6.45), Inches(11.6), Inches(0.4),
                "All four share the same root cause: picking the layer that feels easy to author rather than the one that fits the task.",
                size=12, bold=True, color=WARN)


def slide_recommendation(prs):
    s = add_blank_slide(prs); set_slide_bg(s)
    add_header_bar(s, "Recommendation",
                   "What this means for how the team writes things going forward.",
                   slide_num=12)
    add_textbox(s, Inches(0.55), Inches(1.45), Inches(12.2), Inches(0.4),
                "Default rule",
                size=16, bold=True, color=ACCENT_DARK)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(0.55), Inches(1.85), Inches(12.2), Inches(0.65))
    box.fill.solid(); box.fill.fore_color.rgb = COLOR_SCRIPT_BG
    box.line.color.rgb = COLOR_SCRIPT; box.line.width = Pt(1.25)
    add_textbox(s, Inches(0.85), Inches(1.95), Inches(11.6), Inches(0.5),
                "If you can write a script, write the script. Promote to a skill only when judgment + composition is required. Push to RAG only when corpus size or volatility forces it.",
                size=14, bold=True, color=INK)
    add_textbox(s, Inches(0.55), Inches(2.75), Inches(12.2), Inches(0.4),
                "Three rules of thumb that follow from it",
                size=16, bold=True, color=ACCENT_DARK)
    add_bullets(s, Inches(0.7), Inches(3.15), Inches(12.0), Inches(2.0), [
        ("Fewer skills, sharper skills. Every skill should answer \"what judgment am I teaching the agent here?\" — if the answer is a sequence of steps, write a script.", 0),
        ("Treat skills like code. Version with the repo, link to live paths, add a CI test that fails when those paths move.", 0),
        ("RAG only when no API exists. If the data lives in a DB / catalog / graph, query it directly; RAG is the option of last resort for unstructured corpora.", 0),
    ], size=13)
    add_textbox(s, Inches(0.55), Inches(5.4), Inches(12.2), Inches(0.4),
                "What this means for our work",
                size=16, bold=True, color=ACCENT_DARK)
    add_bullets(s, Inches(0.7), Inches(5.8), Inches(12.0), Inches(1.5), [
        ("Audit existing skills against the rule above — most will demote to scripts or be deleted.", 0),
        ("Keep growing the script library; it is the cheapest, most durable knowledge surface.", 0),
        ("Reserve RAG for live, queryable, evolving corpora — do not use it as a staleness fix for skills.", 0),
    ], size=13)


# ── Build ─────────────────────────────────────────────────────────────────────
def build(out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_title(prs)
    slide_problem(prs)
    slide_at_a_glance(prs)
    slide_script_deepdive(prs)
    slide_skill_deepdive(prs)
    slide_rag_deepdive(prs)
    slide_decision_matrix(prs)
    slide_staleness(prs)
    slide_probabilistic(prs)
    slide_layered(prs)
    slide_antipatterns(prs)
    slide_recommendation(prs)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f"  Wrote {out_path}  ({out_path.stat().st_size / 1024:.1f} KB, {len(prs.slides)} slides)")


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("docs/agent_knowledge_strategy.pptx")
    build(out)
